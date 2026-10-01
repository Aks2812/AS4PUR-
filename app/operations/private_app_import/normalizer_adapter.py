"""
Thin in-memory adapter around reference_scripts/netskope_xlsx_normalize.py.

Reuses its exact row-level parsing/normalization functions
(normalize_hosts, normalize_ports, sanitize_app_name, MAX_HOSTS_PER_APP) -
CLAUDE.md Section 1/8: wrap the proven logic, don't rebuild it. Only the
"read from/write to files on disk" shell around those functions is
replaced: the reference script's own main() writes a JSONL manifest and a
skipped-rows log file, since it's a CLI tool; this adapter returns
structured Python objects instead, since the web app needs them for an
in-memory review page and a background job, not files to be read back
in by a separate step.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from reference_scripts.netskope_xlsx_normalize import (
    MAX_HOSTS_PER_APP,
    normalize_hosts,
    normalize_ports,
    sanitize_app_name,
)

REQUIRED_COLUMNS = ["No", "App Name", "Hosts (IP/CIDR)", "Ports (Protocol/Port)"]


class NormalizerError(Exception):
    """A fatal problem with the file itself (missing columns, unreadable
    workbook) - mirrors the reference script's own fatal-exit cases."""


@dataclass
class ManifestRow:
    row_no: int
    app_name: str
    hosts: list[str]
    protocols: list[dict]
    warnings: list[str] = field(default_factory=list)
    exclusion_reason: str | None = None
    # Set later by service.apply_collision_check(), not by parsing alone:
    # WILL_CREATE | SKIPPED_EXISTS | VALIDATION_EXCLUDED
    status: str = "PENDING"
    # Exactly one of {collision_host, collision_name} is set when status ==
    # SKIPPED_EXISTS - which one distinguishes a destination+port collision
    # from a name collision (CLAUDE.md Section 9: sibling cases, same
    # SKIPPED_EXISTS treatment, different human-readable explanation).
    # collision_port is only set alongside collision_host - the specific
    # existing port range this row's own ports overlapped (CLAUDE.md
    # Section 9: port-level access segmentation - a same-destination row
    # with genuinely disjoint ports is NOT a collision, so collision_host
    # alone no longer implies a collision the way it used to).
    collision_host: str | None = None
    collision_port: str | None = None
    collision_name: str | None = None

    @property
    def is_valid(self) -> bool:
        return self.exclusion_reason is None

    def build_payload(self, publisher_ids: list[int]) -> dict:
        return {
            "app_name": self.app_name,
            "host": ",".join(self.hosts),
            "protocols": self.protocols,
            "publishers": [{"publisher_id": pid} for pid in publisher_ids],
            "use_publisher_dns": False,
            "clientless_access": False,
            "is_user_portal_app": False,
            "trust_self_signed_certs": True,
        }


def parse_workbook(xlsx_path: Path) -> list[ManifestRow]:
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active
    except Exception as exc:
        raise NormalizerError(f"Could not open the uploaded file: {exc}") from exc

    headers = [c.value for c in ws[1]]
    missing = [h for h in REQUIRED_COLUMNS if h not in headers]
    if missing:
        found = ", ".join(str(h) for h in headers if h is not None)
        raise NormalizerError(f"Missing required column(s): {', '.join(missing)}. Found columns: {found}")
    col = {h: headers.index(h) for h in headers if h is not None}

    rows: list[ManifestRow] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        row_no = raw_row[col["No"]]
        if row_no is None:
            continue  # blank trailing row, same as the reference script

        raw_name = raw_row[col["App Name"]]
        raw_hosts = raw_row[col["Hosts (IP/CIDR)"]]
        raw_ports = raw_row[col["Ports (Protocol/Port)"]]

        warnings: list[str] = []
        hosts = normalize_hosts(raw_hosts, warnings, row_no)
        protocols = normalize_ports(raw_ports, warnings, row_no)
        app_name = sanitize_app_name(raw_name if raw_name else f"Unnamed_Row{row_no}", row_no, warnings)

        reasons = []
        if not app_name.strip():
            reasons.append("empty app_name after sanitization")
        if not hosts:
            reasons.append("no valid destination host after normalization")
        if len(hosts) > MAX_HOSTS_PER_APP:
            reasons.append(f"{len(hosts)} hosts exceeds Netskope's {MAX_HOSTS_PER_APP}-host limit per app")
        if not protocols:
            reasons.append("no valid protocol/port after normalization")

        rows.append(
            ManifestRow(
                row_no=row_no,
                app_name=app_name,
                hosts=hosts,
                protocols=protocols,
                warnings=warnings,
                exclusion_reason="; ".join(reasons) if reasons else None,
            )
        )

    return rows
